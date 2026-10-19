from pathlib import Path

from docx import Document


downloads = Path.home() / "Downloads"
files = [
    downloads / "Group2_Technical_Proposal_draft1.docx",
    next(downloads.glob("RE_ GSWA HyLogger Capstone*OFFICIAL*.docx")),
]

for path in files:
    print("\n" + "=" * 90)
    print(path.name)
    print("=" * 90)
    doc = Document(path)
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            print(f"P{i}: {text}")
    for table_index, table in enumerate(doc.tables):
        print(f"\n[TABLE {table_index}]")
        for row in table.rows:
            print(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
