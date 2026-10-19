from pathlib import Path
import sys

from docx import Document


def extract(path_text: str) -> None:
    path = Path(path_text)
    document = Document(path)
    print(f"DOCUMENT: {path}")
    print(f"PARAGRAPHS {len(document.paragraphs)} TABLES {len(document.tables)}")
    for index, paragraph in enumerate(document.paragraphs):
        if not paragraph.text.strip():
            continue
        style_name = paragraph.style.name if paragraph.style is not None else "NoStyle"
        print(f"P{index:03d} [{style_name}] {paragraph.text}")
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            cells = " || ".join(cell.text.replace("\n", " / ") for cell in row.cells)
            print(f"T{table_index:02d}R{row_index:02d} | {cells}")


for source in sys.argv[1:]:
    extract(source)
