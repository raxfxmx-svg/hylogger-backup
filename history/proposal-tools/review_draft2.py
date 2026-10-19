from docx import Document
from pathlib import Path


path = Path(r"C:\path\to\project\Group2_Technical_Proposal_draft2.docx")
doc = Document(path)

print("FILE", path)
print("PARAGRAPHS", len(doc.paragraphs), "TABLES", len(doc.tables))
body = []
for index, paragraph in enumerate(doc.paragraphs):
    value = paragraph.text.strip()
    if value:
        body.append(value)
        print(f"P{index} [{paragraph.style.name}]: {value}")

for table_index, table in enumerate(doc.tables):
    print(f"\n[TABLE {table_index}] rows={len(table.rows)} cols={len(table.columns)}")
    for row_index, row in enumerate(table.rows):
        values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        print(f"R{row_index}: " + " | ".join(values))

print("\nWORD_COUNT_APPROX", len(" ".join(body).split()))
reference_start = next(
    (index for index, text in enumerate(body) if text.upper().replace(" ", "").startswith("5.REFERENCES")),
    len(body),
)
body_without_references = body[:reference_start]
body_without_captions = [
    text for text in body_without_references if not text.lstrip().lower().startswith("figure ")
]
front_matter = 4
print("WORDS_EXCLUDING_REFERENCES", len(" ".join(body_without_references).split()))
print("WORDS_EXCLUDING_REFERENCES_AND_CAPTIONS", len(" ".join(body_without_captions).split()))
print(
    "WORDS_EXCLUDING_FRONT_MATTER_REFERENCES_CAPTIONS",
    len(" ".join(body_without_captions[front_matter:]).split()),
)
print("CORE_TITLE", doc.core_properties.title)
print("CORE_AUTHOR", doc.core_properties.author)
