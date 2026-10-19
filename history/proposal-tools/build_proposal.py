from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\path\to\project\Group2_HyLogger_Lens_Technical_Proposal.docx")

# narrative_proposal preset resolved into explicit tokens
PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_START_END_DXA = 120

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6B76"
LIGHT_FILL = "F4F6F9"
LIGHT_BLUE = "E8EEF5"
WHITE = "FFFFFF"
BLACK = "111111"
BORDER = "C9D2DC"

assessed_fragments: list[str] = []


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, CELL_TOP_BOTTOM_DXA, CELL_START_END_DXA, CELL_TOP_BOTTOM_DXA, CELL_START_END_DXA)


def set_run_font(run, size=None, color=BLACK, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_body(doc, text: str, *, count=True, keep=False, after=8, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style="Normal")
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.keep_together = keep
    p.add_run(text)
    if count:
        assessed_fragments.append(text)
    return p


def add_heading(doc, text: str, level=1, *, count=True):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    if count:
        assessed_fragments.append(text)
    return p


def style_table(table, *, header=True, font_size=9.1, numeric_cols=()):
    set_table_borders(table)
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if header and r == 0:
                set_cell_shading(cell, LIGHT_FILL)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.08
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=font_size, color=DARK_BLUE if (header and r == 0) else BLACK, bold=(header and r == 0))


def add_table_caption(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=9, color=MUTED, bold=True)
    return p


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run = paragraph.add_run()
    set_run_font(begin_run, size=8.5, color=MUTED)
    begin_run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run = paragraph.add_run()
    set_run_font(instr_run, size=8.5, color=MUTED)
    instr_run._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run = paragraph.add_run()
    set_run_font(separate_run, size=8.5, color=MUTED)
    separate_run._r.append(separate)
    placeholder_run = paragraph.add_run("1")
    set_run_font(placeholder_run, size=8.5, color=MUTED)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run = paragraph.add_run()
    set_run_font(end_run, size=8.5, color=MUTED)
    end_run._r.append(end)


doc = Document()
doc.settings.odd_and_even_pages_header_footer = False
section = doc.sections[0]
section.different_first_page_header_footer = False
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# Global styles: narrative_proposal
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333

heading_tokens = {
    1: (16, BLUE, 18, 10),
    2: (13, BLUE, 12, 6),
    3: (12, DARK_BLUE, 8, 4),
}
for level, (size, color, before, after) in heading_tokens.items():
    style = doc.styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True

# Footer only: the title block already carries the document identity.
section.header.paragraphs[0].text = ""
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_page_number(footer)

# Proposal centerpiece opening
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(5)
set_run_font(p.add_run("GROUP 2"), size=11, color=BLUE, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(5)
set_run_font(p.add_run("HyLogger Lens"), size=27, color=NAVY, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(17)
set_run_font(p.add_run("An Interactive Multi-Sensor Drill-Core Explorer"), size=13.5, color=MUTED, italic=True)

meta = doc.add_table(rows=4, cols=2)
meta_rows = [
    ("Client", "Department of Mines, Petroleum and Exploration, Western Australia"),
    ("Project stream", "HyLogger data visualisation"),
    ("Pilot dataset", "05GJD001 (Lake Magenta; 150-316 m)"),
    ("Submission", "Group Proposal | 7 August 2026"),
]
for i, (label, value) in enumerate(meta_rows):
    meta.cell(i, 0).text = label
    meta.cell(i, 1).text = value
set_table_geometry(meta, [2400, 6960])
set_table_borders(meta, color="D7DEE6")
for i, row in enumerate(meta.rows):
    set_cell_shading(row.cells[0], LIGHT_BLUE)
    for j, cell in enumerate(row.cells):
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_run_font(run, size=9.5, color=DARK_BLUE if j == 0 else BLACK, bold=(j == 0))

add_table_caption(doc, "Team")
team = doc.add_table(rows=1, cols=2)
team.cell(0, 0).text = "Name"
team.cell(0, 1).text = "Student number"
for name, sid in [
    ("Contributor 1", "STUDENT_ID_1"),
    ("Contributor 2", "STUDENT_ID_2"),
    ("Contributor 3", "STUDENT_ID_3"),
    ("Contributor 4", "STUDENT_ID_4"),
    ("Contributor 5", "STUDENT_ID_5"),
]:
    cells = team.add_row().cells
    cells[0].text = name
    cells[1].text = sid
set_table_geometry(team, [6500, 2860])
style_table(team, header=True, font_size=9.2, numeric_cols=(1,))

add_heading(doc, "1. Purpose and proposed outcome", 1)
add_body(doc, "By 14 October 2026, Group 2 will deliver a documented, reproducible web prototype that aligns drill-core imagery with depth-indexed VNIR-SWIR and TIR data, lets users compare what the two sensor ranges reveal at the same depth, and exposes quality indicators rather than presenting an unexplained mineral label. The end-to-end pilot will use 05GJD001; the data model and catalogue interface will be designed so additional compatible holes can be added without rewriting the workflow.")
add_body(doc, "The intended users are geologists who need a faster overview and non-specialists who currently face a steep learning curve. The product is not an automated ore/waste classifier. Its purpose is to support exploration, comparison and communication while preserving a clear path back to the underlying spectra, processing fields and source files.")

add_heading(doc, "2. Background and value proposition", 1)
add_body(doc, "HyLogger packages combine high-dimensional spectra, derived mineral parameters, depth information and imagery in specialist formats. At present, a user must find a drill hole in the catalogue, obtain a large ZIP package, understand TSG/BIP/INI files, link it to external metadata, and inspect different sensors using specialist knowledge or custom scripts. This fragmented workflow is slow, difficult to reproduce and easy to misinterpret, particularly when SWIR and TIR respond differently to the same material (Chard, 2026).")
add_body(doc, "The scale makes full ingestion unrealistic for this capstone. The NCI WA directory contains 1,951 ZIP packages whose displayed sizes total approximately 2.384 TB; the accompanying catalogue contains 2,298 records across 29 metadata fields, which are not assumed to map one-to-one to ZIP files. A pilot-first design therefore creates more value than an untested archive-wide pipeline: it proves the data contract, interaction model and validation checks on one known package, then documents the requirements for scaling.")
add_body(doc, "HyLogger Lens will reduce repeated manual handling, make depth and sensor comparisons consistent, and give every chart a reproducible transformation path. Linked views and plain-language guidance should shorten the time from data access to geological questioning. Explicit masks, error/SNR fields and provenance will also reduce false confidence. Because the workflow uses open-source Python components and cached derived files, the proposed cash cost is zero and the prototype can be rerun on UWA or client hardware.")

add_heading(doc, "3. Data baseline and scope", 1)
add_body(doc, "The catalogue identifies 05GJD001 as a Lake Magenta gold-project hole spanning 150-316 m. The downloaded 782.4 MB package contains 33 tray JPEGs, a mosaic, VNIR-SWIR and TIR TSG/BIP/INI/CAL groups, line-scan data and a high-resolution profilometer file. Both spectral groups report 25,938 sample positions; the VNIR-SWIR group has 531 channels and the TIR group 341. These counts form the first validation baseline. Catalogue coordinates and other external metadata will be joined explicitly, because they are not assumed to be embedded in the raw package.")
add_body(doc, "Minimum scope is one complete drill hole with catalogue search, depth navigation, linked imagery, sensor-specific spectra, mineral/parameter tracks, quality overlays and export of the current selection. Stretch scope is up to four additional compatible holes and cross-hole comparison. Geochemistry, assays, 3D desurveying and supervised mineralisation prediction are deferred until the client supplies suitable data and definitions.")

add_heading(doc, "4. Method", 1)
add_heading(doc, "4.1 Ingestion, validation and provenance", 2)
add_body(doc, "A pinned Python environment will use pytsg to test package parsing and nvcl-kit only where it improves catalogue access. Raw files remain immutable and outside Git. A manifest will record filename, size, checksum, parser version and extraction time. Validation will test required files, sample counts, sensor/depth alignment, monotonic depth, duplicate or missing depths, finite spectra, tray references and the presence of masks and quality variables. Validated metadata will be stored in SQLite and analysis-ready arrays in Parquet or another benchmarked columnar form.")

add_heading(doc, "4.2 Exploratory analysis and technique selection", 2)
add_body(doc, "The team will first visualise representative raw and preprocessed spectra, depth distributions, missingness, error/SNR values and mineral parameters. At least three exploratory approaches will then be compared: (1) principal component analysis for compact spectral structure; (2) k-means or hierarchical clustering for recurring spectral intervals; and (3) change-point detection or rolling-distance measures for transitions down-hole. Alternatives will be evaluated by stability under reasonable preprocessing choices, ability to reproduce the same result from code, and interpretability when displayed beside tray imagery. Sensor outputs will be compared at matched depths, but disagreement will be shown rather than collapsed into a single 'truth'.")

add_heading(doc, "4.3 Interaction and uncertainty communication", 2)
add_body(doc, "A Streamlit/Plotly interface will link four views: catalogue and hole summary; tray image with current depth; SWIR and TIR spectra at that depth; and depth tracks for selected parameters or mineral interpretations. Brushing, zooming and a shared depth cursor will keep the views synchronised. Users can toggle sensors, raw/processed views and quality layers, inspect definitions through tooltips, and export the visible subset with provenance. Error, SNR, NIL status and final-mask fields will be displayed as data-quality evidence. File suffixes or catalogue process levels will not be converted into confidence categories without documentation and client confirmation.")

add_heading(doc, "4.4 Evaluation and collaborative delivery", 2)
add_body(doc, "Technical tests will cover parsers, depth joins and exported values, with spot checks back to the source files. Usability evaluation will involve at least five non-specialist participants completing five tasks: find the pilot hole, move to a requested depth, compare SWIR/TIR, identify a masked or low-quality interval, and export a selection. Completion, errors, time and comments will be recorded; one client or domain-expert walkthrough will be requested to check terminology and interpretation. Work will be managed in a shared Git repository using issues, feature branches, pull requests, a second-person review and weekly tagged releases. Large data and generated caches will remain outside version control, while configuration, tests and environment locks make the workflow reproducible.")

add_heading(doc, "5. Deliverables and acceptance criteria", 1)
add_table_caption(doc, "Table 1. SMART deliverables")
deliverables = doc.add_table(rows=1, cols=4)
for j, text in enumerate(("ID", "Deliverable", "Acceptance criterion", "Due")):
    deliverables.cell(0, j).text = text
deliverable_rows = [
    ("D1", "Data contract and validated pilot extract", "05GJD001 parsed from immutable source; all 25,938 positions reconciled or exceptions logged; manifest, schema and catalogue join documented.", "30 Aug"),
    ("D2", "Exploratory analysis report", "At least three techniques compared using stated stability, reproducibility and interpretability criteria; SWIR/TIR differences and limitations illustrated.", "13 Sep"),
    ("D3", "Interactive MVP", "Linked depth, tray-image, spectra and parameter views; sensor/quality controls; filtered export with provenance; reproducible start-up instructions.", "27 Sep"),
    ("D4", "Evaluation and refined release", "Five users complete five tasks; results and resulting interface changes documented; client/domain walkthrough requested and recorded if available.", "11 Oct"),
    ("D5", "Final package", "Versioned code, tests, data dictionary, quick-start guide, demonstration, pitch assets and final group report delivered.", "14 Oct"),
]
for row_data in deliverable_rows:
    cells = deliverables.add_row().cells
    for j, value in enumerate(row_data):
        cells[j].text = value
set_table_geometry(deliverables, [560, 2050, 5400, 1350])
style_table(deliverables, header=True, font_size=8.55, numeric_cols=(0, 3))

add_heading(doc, "6. Timeline, responsibilities and cost", 1)
add_table_caption(doc, "Table 2. Delivery schedule")
timeline = doc.add_table(rows=1, cols=3)
for j, text in enumerate(("Period", "Team focus", "Evidence / milestone")):
    timeline.cell(0, j).text = text
timeline_rows = [
    ("10-16 Aug", "Confirm research questions, create repository and data contract; review catalogue and raw package.", "Requirements and risk log approved by team."),
    ("17-30 Aug", "Build parser, manifest, validation tests and cached pilot extract.", "D1 demonstration and peer review."),
    ("31 Aug-13 Sep", "Run exploratory methods; compare SWIR/TIR; prototype linked charts.", "D2 findings and chosen interaction design."),
    ("14-27 Sep", "Integrate MVP; add definitions, quality overlays and export.", "D3 feature-complete release."),
    ("28 Sep-11 Oct", "Usability tasks, client walkthrough, fixes, documentation and release freeze.", "D4 evidence and final candidate."),
    ("12-14 Oct", "Mini-conference, pitch video and final report/code handover.", "D5 submitted by 14 October."),
]
for row_data in timeline_rows:
    cells = timeline.add_row().cells
    for j, value in enumerate(row_data):
        cells[j].text = value
set_table_geometry(timeline, [1700, 4490, 3170])
style_table(timeline, header=True, font_size=8.7)

add_body(doc, "The five members will work in paired streams (data/analysis and interface/evaluation), with an integration owner rotating each sprint so knowledge is shared and no component has a single maintainer. The estimated in-kind effort is approximately 400 person-hours (five members x eight hours per week x ten weeks).", count=True)

add_table_caption(doc, "Table 3. Cost estimate")
costs = doc.add_table(rows=1, cols=3)
for j, text in enumerate(("Item", "Basis", "Cash cost (AUD)")):
    costs.cell(0, j).text = text
cost_rows = [
    ("Data and software", "Open-file NCI data; Python, pytsg, Streamlit, Plotly and Git", "$0"),
    ("Compute and storage", "Pilot processed on existing student/UWA hardware; cached outputs kept small", "$0"),
    ("Hosting", "Local reproducible demo; free-tier hosted summary only if feasible", "$0"),
    ("Labour", "Approx. 400 student hours contributed in kind", "$0 cash"),
]
for row_data in cost_rows:
    cells = costs.add_row().cells
    for j, value in enumerate(row_data):
        cells[j].text = value
set_table_geometry(costs, [2450, 5110, 1800])
style_table(costs, header=True, font_size=8.8, numeric_cols=(2,))

add_heading(doc, "7. Scope controls and expected impact", 1)
add_body(doc, "The main risks are unfamiliar binary formats, ambiguous domain labels, package heterogeneity and the temptation to scale before validating the pilot. These are controlled through immutable raw data, explicit schemas and checksums, client questions recorded as issues, sensor-specific displays, and a hard MVP freeze on 27 September. If parser or metadata gaps remain, the team will deliver the validated subset and document unsupported fields rather than silently dropping them. Full-archive ingestion, paid cloud infrastructure and unvalidated classification remain out of scope.")
add_body(doc, "Success will be a working, inspectable pathway from one catalogue record to aligned evidence at depth: image, spectra, interpreted parameters, quality information and export. That outcome is small enough to complete, useful enough to demonstrate to GSWA and industry, and structured so later teams can add holes, geochemistry or 3D context without discarding the foundation.")

add_heading(doc, "References", 1, count=False)
references = [
    "Chard, J. (2026). GSWA HyLogger Capstone - First Meeting & Dataset Access. Client meeting transcript, 3 August 2026.",
    "Geological Survey of Western Australia (GSWA). (2026a). gswa-hylogger-2026_07_06.xlsx. Client-supplied HyLogger catalogue.",
    "Geological Survey of Western Australia (GSWA). (2026b). 05GJD001 HyLogger package. Client-supplied data extract.",
    "NCI Australia. (2026). THREDDS catalogue: rs07/WA. https://thredds.nci.org.au/thredds/catalog/rs07/WA/catalog.html",
    "pytsg. (2025). Python package for reading TSG file packages. https://pypi.org/project/pytsg/",
    "NVCL Kit. (2026). Python tools for National Virtual Core Library data access. https://pypi.org/project/nvcl-kit/",
]
for ref in references:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(ref)
    set_run_font(run, size=9.3, color=BLACK)

# Estimated assessed word count excludes tables and references, matching the brief.
word_count = len(re.findall(r"\b[\w'-]+\b", " ".join(assessed_fragments)))
count_note = doc.paragraphs[10] if len(doc.paragraphs) > 10 else None

# Core metadata and accessibility properties
doc.core_properties.title = "HyLogger Lens - Group 2 Technical Proposal"
doc.core_properties.subject = "UWA Data Science Capstone group proposal"
doc.core_properties.author = "Group 2"
doc.core_properties.keywords = "HyLogger, hyperspectral, SWIR, TIR, drill core, visualisation"
doc.core_properties.comments = f"Assessed prose word count (tables and references excluded): {word_count}"

doc.save(OUTPUT)
print(f"Saved: {OUTPUT}")
print(f"Assessed prose word count (tables and references excluded): {word_count}")
