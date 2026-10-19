from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document
from docx.oxml.ns import qn


path = Path('C:\\path\\to\\project\\Group2_HyLogger_Lens_Group_Proposal_Review.docx')
doc = Document(path)
section = doc.sections[0]

checks = {
    "page_width_letter": section.page_width == 7772400,
    "page_height_letter": section.page_height == 10058400,
    "margin_top_1in": section.top_margin == 914400,
    "margin_right_1in": section.right_margin == 914400,
    "margin_bottom_1in": section.bottom_margin == 914400,
    "margin_left_1in": section.left_margin == 914400,
    "table_count": len(doc.tables) == 8,
    "heading_count": len([p for p in doc.paragraphs if p.style.name.startswith("Heading")]) == 20,
}

for table_index, table in enumerate(doc.tables, start=1):
    grid_widths = [int(c.get(qn("w:w"))) for c in table._tbl.tblGrid]
    checks[f"table_{table_index}_grid_9360"] = sum(grid_widths) == 9360
    checks[f"table_{table_index}_no_fixed_row_height"] = all(
        row._tr.get_or_add_trPr().find(qn("w:trHeight")) is None for row in table.rows
    )
    for row_index, row in enumerate(table.rows, start=1):
        cell_widths = [
            int(cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w")))
            for cell in row.cells
        ]
        checks[f"table_{table_index}_row_{row_index}_matches_grid"] = cell_widths == grid_widths

with ZipFile(path) as package:
    all_xml = "\n".join(
        package.read(name).decode("utf-8", errors="ignore")
        for name in package.namelist()
        if name.endswith(".xml")
    )
    checks["page_field_present"] = " PAGE " in all_xml
    checks["no_internal_tool_tokens"] = not re.search(
        r"turn\d+(search|fetch)|codex-file-citation|:codex-", all_xml
    )
    checks["no_tracked_changes"] = not re.search(r"<w:(ins|del)(?:\s|>)", all_xml)

failed = [name for name, passed in checks.items() if not passed]
print(f"File: {path}")
print(f"Size: {path.stat().st_size} bytes")
print(f"Checks passed: {len(checks) - len(failed)}/{len(checks)}")
if failed:
    print("FAILED")
    for name in failed:
        print(name)
    raise SystemExit(1)
print("All structural checks passed.")
