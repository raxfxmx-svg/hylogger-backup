from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document


path = Path('C:\\Users\\USER\\Desktop\\Individual_Report_Draft_3.docx')
doc = Document(path)

print("FILE", path)
print("PARAGRAPHS", len(doc.paragraphs), "TABLES", len(doc.tables))
texts = []
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        texts.append(text)
        print(f"P{index} [{paragraph.style.name}]: {text}")

for table_index, table in enumerate(doc.tables):
    print(f"\n[TABLE {table_index}] rows={len(table.rows)} cols={len(table.columns)}")
    for row_index, row in enumerate(table.rows):
        cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        print(f"R{row_index}: " + " | ".join(cells))

all_text = " ".join(texts)
words = re.findall(r"\b[\w'-]+\b", all_text, flags=re.UNICODE)
print("\nWORD_COUNT_APPROX", len(words))
print("CORE_AUTHOR", doc.core_properties.author)
print("CORE_TITLE", doc.core_properties.title)

with ZipFile(path) as package:
    names = set(package.namelist())
    xml = "\n".join(
        package.read(name).decode("utf-8", errors="ignore")
        for name in names
        if name.endswith(".xml")
    )
    print("COMMENTS_PRESENT", "word/comments.xml" in names)
    print("TRACKED_CHANGES_PRESENT", bool(re.search(r"<w:(ins|del)(?:\s|>)", xml)))
