from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path('C:\\path\\to\\project\\Group2_HyLogger_Lens_Group_Proposal_Review.docx')

# narrative_proposal preset, with a named East-Asian font override.
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6B76"
LIGHT_FILL = "F4F6F9"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "DCEAF7"
WHITE = "FFFFFF"
BLACK = "111111"
BORDER = "C9D2DC"
GREEN = "2F6B55"


def set_run_font(run, size=None, color=BLACK, bold=None, italic=None, latin="Calibri", east_asia="Microsoft YaHei"):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), latin)
    rpr.rFonts.set(qn("w:hAnsi"), latin)
    rpr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table column widths must sum to  {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def style_table(table, header=True, font_size=8.9, centered_cols=()):
    set_table_borders(table)
    if header:
        mark_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            if header and row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2.5)
                paragraph.paragraph_format.line_spacing = 1.08
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index in centered_cols else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        color=DARK_BLUE if (header and row_index == 0) else BLACK,
                        bold=(header and row_index == 0),
                    )


def add_page_number(paragraph):
    set_run_font(paragraph.add_run('Page  '), size=8.5, color=MUTED)
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run._r.append(instr)
    sep_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(separate)
    set_run_font(paragraph.add_run("1"), size=8.5, color=MUTED)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    set_run_font(paragraph.add_run(' Page'), size=8.5, color=MUTED)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, after=8, keep=False):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.333
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.keep_together = keep
    set_run_font(paragraph.add_run(text), size=11)
    return paragraph


def add_labeled_paragraph(doc, label, text, after=6):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    set_run_font(paragraph.add_run(f"{label}："), size=10.8, color=DARK_BLUE, bold=True)
    set_run_font(paragraph.add_run(text), size=10.8)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), size=9, color=MUTED, bold=True)
    return paragraph


def add_reference(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.28)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    set_run_font(paragraph.add_run(text), size=9.1)


doc = Document()
doc.settings.odd_and_even_pages_header_footer = False
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
section.different_first_page_header_footer = False

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333

heading_tokens = {
    1: (16, BLUE, 18, 10),
    2: (13, BLUE, 12, 6),
    3: (11.5, DARK_BLUE, 8, 4),
}
for level, (size, color, before, after) in heading_tokens.items():
    style = doc.styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True

section.header.paragraphs[0].text = ""
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_page_number(footer)

# proposal_centerpiece title block
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(5)
set_run_font(p.add_run('UWA Data Science Capstone - Group 2'), size=11, color=BLUE, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
set_run_font(p.add_run("HyLogger Lens"), size=27, color=NAVY, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
set_run_font(p.add_run('Interactive analysis and visualisation of multi-sensor drill-core data'), size=14, color=MUTED, bold=True)

meta = doc.add_table(rows=4, cols=2)
meta_rows = [
    ('Client', 'Department of Mines, Petroleum and Exploration / Geological Survey of Western Australia (DMPE / GSWA)'),
    ('Project direction', 'HyLogger data visualisation'),
    ('End-to-end pilot', '05GJD001 (150-316 m; VNIR-SWIR, TIR and tray images)'),
    ('Draft status', 'Internal review draft mapped to the marking criteria; final submission limited to 1,500 words'),
]
for row_index, (label, value) in enumerate(meta_rows):
    meta.cell(row_index, 0).text = label
    meta.cell(row_index, 1).text = value
set_table_geometry(meta, [2200, 7160])
set_table_borders(meta, color="D7DEE6")
for row in meta.rows:
    set_cell_shading(row.cells[0], LIGHT_BLUE)
    for col_index, cell in enumerate(row.cells):
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                set_run_font(run, size=9.3, color=DARK_BLUE if col_index == 0 else BLACK, bold=(col_index == 0))

add_caption(doc, 'Team members')
team = doc.add_table(rows=1, cols=2)
team.cell(0, 0).text = 'Name'
team.cell(0, 1).text = 'Student ID'
for name, student_number in [
    ("Contributor 1", "STUDENT_ID_1"),
    ("Contributor 2", "STUDENT_ID_2"),
    ("Contributor 3", "STUDENT_ID_3"),
    ("Contributor 4", "STUDENT_ID_4"),
    ("Contributor 5", "STUDENT_ID_5"),
]:
    cells = team.add_row().cells
    cells[0].text = name
    cells[1].text = student_number
set_table_geometry(team, [6500, 2860])
style_table(team, header=True, font_size=9.1, centered_cols=(1,))

note = doc.add_table(rows=1, cols=1)
note.cell(0, 0).text = 'Review note: this draft follows the assessment rubric. The appendix mapping is for internal review and may be omitted from the submitted English version.'
set_table_geometry(note, [9360])
set_table_borders(note, color="B8CADB")
set_cell_shading(note.cell(0, 0), LIGHT_FILL)
for run in note.cell(0, 0).paragraphs[0].runs:
    set_run_font(run, size=9.2, color=DARK_BLUE, bold=True)

add_heading(doc, '1. Project title', 1)
add_body(doc, 'HyLogger Lens: interactive analysis and visualisation of multi-sensor drill-core data.')

add_heading(doc, '2. Purpose', 1)
add_body(doc, 'By October 14, 2026, the team will deliver a reproducible Python/Streamlit web prototype using 05GJD001 as the end-to-end pilot. Catalog metadata, core-tray images, VNIR-SWIR/TIR spectra, mineral/parameter tracks, and quality fields share a depth axis. Users can find boreholes and depths, compare sensors, inspect quality evidence, and export selected data with provenance.')
add_body(doc, 'The minimum scope fully supports one borehole. Add up to four more and cross-borehole comparison if formats and time permit. Unvalidated spectra will not be classified directly as ore or waste, and machine-learning classification is not a required deliverable. This supports intuitive interaction while keeping interpretation risk manageable within the course schedule.')

add_heading(doc, '3. Background and value proposition (2 marks)', 1)
add_heading(doc, '3.1 Understanding the data and current interpretation practices', 2)
add_body(doc, 'The NCI WA directory lists 1,951 ZIP packages totalling approximately 2.384 TB by displayed size. The client catalogue contains 2,298 records and 29 metadata fields, so records and ZIPs are not necessarily one-to-one. Packages typically contain TSG, BIP, INI, CAL and core images. Users locate boreholes, download packages, parse them with The Spectral Geologist or scripts, then obtain coordinates, surveys, geological logs and assays from the catalogue, WAMEX/WAPIMS or other databases.')
add_body(doc, 'VNIR-SWIR and TIR respond to different minerals and materials. The client asked for explanations of differences at the same depth. Preserve sensor provenance, original parameters and quality fields, using core images as geological context.')

add_heading(doc, '3.2 Expected challenges in interpreting client data', 2)
add_labeled_paragraph(doc, 'Scale challenges', 'The full archive is approximately 2.384 TB, with individual files exceeding 13 GB. Downloading and parsing the entire archive is outside the course schedule.')
add_labeled_paragraph(doc, 'Format and skills barriers', 'TSG/BIP structures require specialist software or parsers. High-dimensional spectra are difficult to interpret in ordinary spreadsheets, creating a barrier for non-specialists.')
add_labeled_paragraph(doc, 'Multi-source alignment challenges', 'Join images, both sensor groups, catalogue coordinates and future geochemistry/logging data by borehole and depth. Incorrect joins can produce more misleading conclusions than missing data.')
add_labeled_paragraph(doc, 'Interpretation risks', 'The interface may hide sensor differences, missing values, masks, errors, SNR and NIL states. Filename suffixes and catalogue Process Level values must not become confidence ratings without documentation and client confirmation.')

add_heading(doc, '3.3 Value of data science methods', 2)
add_labeled_paragraph(doc, 'Efficiency', 'Cache parsed data locally for analysis and link views by depth to reduce repeated file opening and manual matching. Measure task completion time during user testing.')
add_labeled_paragraph(doc, 'Consistency', 'A shared schema and automated quality checks apply the same file, sample-count, depth-order, duplicate, missing-value, and sensor-alignment checks to each borehole.')
add_labeled_paragraph(doc, 'Reproducibility', 'Keep raw data read-only. Manifests record file sizes, SHA-256, parser versions, and transformation times. Code, environment lockfiles, and parameter settings allow another member to reproduce the results.')
add_labeled_paragraph(doc, 'Cost', 'Open data, Python, pytsg, Streamlit, Plotly, Git and existing UWA/student equipment give an estimated pilot cash cost of AUD 0. Full-archive cloud storage and paid deployment are outside scope.')

add_heading(doc, '4. Deliverables, schedule and costs (3 marks)', 1)
add_heading(doc, '4.1 Expected outcomes: SMART deliverables', 2)
add_caption(doc, 'Table 1. Measurable, achievable, relevant and time-bound deliverables')
deliverables = doc.add_table(rows=1, cols=4)
for col, text in enumerate(('ID', 'Deliverable', 'Measurable acceptance criteria', 'Deadline')):
    deliverables.cell(0, col).text = text
deliverable_rows = [
    ("D1", 'Data queries, dictionary, and validated pilot extraction', 'Complete the 05GJD001 file inventory and catalogue join. Reconcile all 25,938 sample positions or record exceptions; trace 33 tray images, 531 VNIR-SWIR channels and 341 TIR channels to their sources.', '23 August'),
    ("D2", 'Data exploration and method comparison report', 'Test at least 3 analytical methods for stability, reproducibility, interpretability and runtime cost; check 20 randomly selected depths against the source data.', 'September 6'),
    ("D3", 'Interactive visualisation MVP', 'Implement four linked views: catalogue/borehole overview, tray images, SWIR/TIR spectra and depth tracks. Support quality layers, parameter controls and provenance in exports. After preloading, at least 90% of 30 typical interactions should update within 2 seconds.', '20 September'),
    ("D4", 'Internal testing and revised prototype', 'Ask at least 5 non-specialist users to complete 5 tasks each. Record completion, time, errors and feedback; target at least 80% completion without prompts. Request and record one client or domain-expert walkthrough.', '4 October'),
    ("D5", 'Final demonstration, report, and code package', 'Deliver tagged code, automated tests, a data dictionary, quick-start instructions, a demonstration, pitch materials and the final group report. Document known limitations in the README and report.', 'October 14'),
]
for row_data in deliverable_rows:
    cells = deliverables.add_row().cells
    for col, value in enumerate(row_data):
        cells[col].text = value
set_table_geometry(deliverables, [560, 1900, 5520, 1380])
style_table(deliverables, header=True, font_size=8.2, centered_cols=(0, 3))

add_heading(doc, '4.2 Analysis and visualisation outputs', 2)
add_body(doc, 'Analysis outputs cover catalog/file scale, fields and missingness, depth/tray coverage, sensor quality distributions, same-depth SWIR/TIR comparison, and PCA/clustering/change-point comparisons. Visual outputs include a filtered borehole overview, depth-linked tray images, paired spectra, mineral/parameter tracks, quality masks, and exports. Every result figure identifies the borehole, depth, sensor, processing parameters, and source.')

add_heading(doc, '4.3 Demonstration, report, and code', 2)
add_body(doc, 'Deliver a live web demonstration or backup recording, final group report, pitch materials and a reproducible repository. Include pinned dependencies, configuration, data manifests, parsing/validation tests, startup instructions, a data dictionary and known limitations. Demonstrate borehole search, depth selection, SWIR/TIR comparison, quality inspection and export.')

gantt_heading = add_heading(doc, '4.4 Completion schedule (Gantt chart)', 2)
gantt_heading.paragraph_format.page_break_before = True
add_caption(doc, 'Figure 1. Schedule from data discovery to the final report; dark blue cells mark planned work')
gantt = doc.add_table(rows=1, cols=7)
gantt_headers = ['Workflow', "8/10-8/23", "8/24-9/6", "9/7-9/20", "9/21-10/4", "10/5-10/11", "10/12-10/14"]
for col, text in enumerate(gantt_headers):
    gantt.cell(0, col).text = text
gantt_rows = [
    ('Confirm requirements, query catalogs, and inventory data', [1, 1, 0, 0, 0, 0]),
    ('Parsing, depth joins and quality validation', [1, 1, 0, 0, 0, 0]),
    ('Exploratory analysis and method comparison', [0, 1, 1, 0, 0, 0]),
    ('Visualisation prototype and feature integration', [0, 1, 1, 1, 0, 0]),
    ('Internal testing, client walkthrough and fixes', [0, 0, 0, 1, 1, 0]),
    ('Documentation, pitch, demonstration and final report', [0, 0, 1, 1, 1, 1]),
]
for label, active_cells in gantt_rows:
    cells = gantt.add_row().cells
    cells[0].text = label
    for index, active in enumerate(active_cells, start=1):
        cells[index].text = "●" if active else ""
        if active:
            set_cell_shading(cells[index], BLUE)
set_table_geometry(gantt, [2400, 1160, 1160, 1160, 1160, 1160, 1160])
style_table(gantt, header=True, font_size=7.45, centered_cols=(1, 2, 3, 4, 5, 6))
for row in gantt.rows[1:]:
    for cell in row.cells[1:]:
        if cell.text == "●":
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9, color=WHITE, bold=True)
add_body(doc, 'D4 internal testing ends on 4 October. Freeze the release on 11 October, attend the Mini-Conference and submit the pitch on 12-13 October, and submit the final group report on 14 October. The 10-day interval after testing covers fixes, verification and writing.', after=6)

add_heading(doc, '4.5 Costs', 2)
add_caption(doc, 'Table 2. Estimated project costs')
costs = doc.add_table(rows=1, cols=3)
for col, text in enumerate(('Project', 'Estimate basis', 'Cash cost (AUD)')):
    costs.cell(0, col).text = text
for row_data in [
    ('Data and software', 'NCI open data; Python, pytsg, Streamlit, Plotly and Git/DVC', "$0"),
    ('Compute and storage', 'Use existing student/UWA equipment; process 05GJD001 and a small number of compatible additional boreholes.', "$0"),
    ('Deployment', 'A reproducible local demonstration is required; free hosting is optional.', "$0"),
    ('Labour', '5 people x approximately 8 hours/week x 10 weeks = approximately 400 course hours', '$0 cash'),
]:
    cells = costs.add_row().cells
    for col, value in enumerate(row_data):
        cells[col].text = value
set_table_geometry(costs, [2200, 5300, 1860])
style_table(costs, header=True, font_size=8.6, centered_cols=(2,))

add_heading(doc, '5. Methods (5 marks)', 1)
add_heading(doc, '5.1 Initial data queries', 2)
add_labeled_paragraph(doc, 'Catalogue query', 'Filter the client workbook and NCI THREDDS catalog by borehole name, Availability, Commodity, Project, Location, coordinates, depth range, and file size. Distinguish catalog records, ZIP packages, and derived/subset data.')
add_labeled_paragraph(doc, 'Pilot selection', '05GJD001 is already downloaded and includes both sensor groups and tray images. Its catalogue coverage is 150-316 m, enabling end-to-end validation without further downloads.')
add_labeled_paragraph(doc, 'File inventory', 'Record relative paths, types, byte counts, and SHA-256 for raw files. Keep raw data read-only and write cleaned or transformed data to a separate processed directory.')
add_labeled_paragraph(doc, 'Parsing and joining', 'Read spectra and parameters with pytsg. Align VNIR-SWIR, TIR, and image indexes by HoleID and Depth, then explicitly join external metadata such as catalog coordinates.')
add_labeled_paragraph(doc, 'Quality baseline', 'Check required files, 25,938 sample positions, monotonic depths, repeated/missing depths, nonfinite spectra, tray references, Final Mask, Error, SNR, and NIL states. Record every exception in the validation report.')

add_heading(doc, '5.2 Testing data exploration and analysis methods', 2)
add_caption(doc, 'Table 3. Exploration and analysis methods to compare')
methods = doc.add_table(rows=1, cols=4)
for col, text in enumerate(('Method', 'Question addressed', 'Test approach', 'Retention criteria')):
    methods.cell(0, col).text = text
method_rows = [
    ('Basic profiles and quality statistics', 'Which depths contain missing data, anomalies or sensor differences?', 'Plot depth coverage, parameter distributions and SNR/Error/Mask fields; check 20 randomly selected depths against the source.', 'Results trace directly to source data; retain this required baseline.'),
    ('Principal component analysis (PCA)', 'Can high-dimensional spectra be reduced to a few interpretable variation axes?', 'Compare explained variance and depth continuity for raw and normalized inputs.', 'Stable under preprocessing, affordable to compute and interpretable alongside images.'),
    ('K-means and hierarchical clustering', 'Can recurring spectral intervals be identified?', 'Compare cluster counts, initializations, and sensor inputs; assess silhouette scores and depth continuity.', 'Clusters are stable, are not driven by masks/missingness, and can be understood by domain experts.'),
    ('Rolling spectral distance and change points', 'Which depths may contain rapid changes or boundaries?', 'Test window sizes and thresholds against tray images and quality fields.', 'Change points should remain stable under reasonable parameters without over-detecting noise.'),
]
for row_data in method_rows:
    cells = methods.add_row().cells
    for col, value in enumerate(row_data):
        cells[col].text = value
set_table_geometry(methods, [1800, 2350, 3000, 2210])
style_table(methods, header=True, font_size=8.2)
add_body(doc, 'Test at least three advanced methods, then retain one or two based on stability, reproducibility, interpretability, and runtime. Keep baseline profiles for comparison. Model outputs identify intervals for investigation rather than expert-validated mineral ground truth.')

add_heading(doc, '5.3 Making the workflow understandable and outputs usable', 2)
add_labeled_paragraph(doc, 'Parameter controls', 'Provide borehole, depth, sensor, raw/processed view, smoothing window, analysis method, and cluster-count controls. Explain defaults, units, and valid ranges in the interface.')
add_labeled_paragraph(doc, 'Integrate domain knowledge', 'Show borehole, tray, depth, sensor, and parameter definitions beside images and spectra. Explain mineral and quality fields using the client glossary/quick reference and record unresolved terminology.')
add_labeled_paragraph(doc, 'Communicate uncertainty', 'Display missing values, Final Mask, Error, SNR, and NIL states as separate tracks or overlays. Show sensor disagreement directly and mark unknown or unvalidated values explicitly.')
add_labeled_paragraph(doc, 'Output usability', 'Users can export the selected depth range as CSV/Parquet with provenance recording borehole, depth, sensor, parameters, parser version, and source files.')

add_heading(doc, '5.4 Visualising the process and results', 2)
add_body(doc, 'Show the workflow from catalogue query through raw manifests, parsing, validation, depth alignment, method testing and caching to the interface. Summarise checks on a validation page. Synchronise four views: borehole overview, tray images and depth cursor, SWIR/TIR spectra at the selected depth, and mineral/parameter/quality tracks. Optional PCA, clustering and change-point layers share the same selection.')
add_body(doc, 'Visualisation acceptance tests five user tasks: find a borehole, locate a depth, compare sensors, identify quality issues, and export results. Legends, labels, units, tooltips, and text states supplement colour encoding.')

add_heading(doc, '5.5 Collaborative tools for data and code version control', 2)
add_labeled_paragraph(doc, 'Code collaboration', 'Use GitHub, Issues/Project boards, short-lived feature branches, and pull requests. Require another member to review each merge and automated tests to pass.')
add_labeled_paragraph(doc, 'Data version', 'Keep large raw files outside Git. Record data versions, SHA-256 hashes and source URLs in DVC metadata or immutable manifests. Version small test fixtures and the data dictionary with the code.')
add_labeled_paragraph(doc, 'Reproducible environment', 'Pin Python and library versions and store analysis parameters in configuration files. Create weekly release tags; D1-D5 correspond to recoverable milestone versions.')
add_labeled_paragraph(doc, 'Team operation', 'Five members work in paired data/analysis and interface/evaluation streams, rotating the integration owner each sprint. Weekly meetings review risks, dependencies, completion criteria, and client questions to distribute knowledge.')

add_heading(doc, 'References', 1)
for reference in [
    'Chard, J. (2026). GSWA HyLogger Capstone - First Meeting & Dataset Access. Client meeting transcript, 3 August 2026.',
    'Geological Survey of Western Australia (2026a). gswa-hylogger-2026_07_06.xlsx. Client-provided HyLogger catalogue.',
    'Geological Survey of Western Australia (2026b). 05GJD001 HyLogger package, client-provided data extraction.',
    "NCI Australia（2026）。THREDDS catalogue: rs07/WA。https://thredds.nci.org.au/thredds/catalog/rs07/WA/catalog.html",
    "pytsg。Python package for reading TSG file packages。https://pypi.org/project/pytsg/",
    "NVCL Kit。Python tools for National Virtual Core Library data access。https://pypi.org/project/nvcl-kit/",
]:
    add_reference(doc, reference)

add_heading(doc, 'Appendix: marking criteria mapping (internal review)', 1)
add_caption(doc, 'Table 4. Mapping this draft to the marking criteria')
rubric = doc.add_table(rows=1, cols=3)
for col, text in enumerate(('Assessment criterion', 'Location in this draft', 'Evidence available to the assessor')):
    rubric.cell(0, col).text = text
rubric_rows = [
    ('Project name and purpose', 'Sections 1-2', 'Define the project name, deadline, minimum scope, extensions, and exclusions.'),
    ('Background and value proposition (2 marks)', 'Sections 3.1-3.3', 'Current practice, four client challenges, efficiency, consistency, reproducibility and cost.'),
    ('Deliverables, schedule and costs (3 marks)', 'Sections 4.1-4.5', 'D1-D5 SMART criteria, analysis and visualisation outputs, demonstration/report/code, Gantt chart, and AUD 0 cost.'),
    ('Methods (5 marks)', 'Sections 5.1-5.5', 'Initial queries, at least three method trials, parameters/domain knowledge/uncertainty, process and result visualisation, and Git/DVC collaboration.'),
]
for row_data in rubric_rows:
    cells = rubric.add_row().cells
    for col, value in enumerate(row_data):
        cells[col].text = value
set_table_geometry(rubric, [1900, 2100, 5360])
style_table(rubric, header=True, font_size=8.6)

doc.core_properties.title = 'HyLogger Lens - Group Proposal Review Draft'
doc.core_properties.subject = 'UWA Data Science Capstone Group Proposal'
doc.core_properties.author = "Group 2"
doc.core_properties.keywords = 'HyLogger, SWIR, TIR, drill core, data visualisation, group proposal'
doc.core_properties.comments = 'Internal review draft; the final English submission is limited to 1,500 words.'

doc.save(OUTPUT)

plain_text = "".join(paragraph.text for paragraph in doc.paragraphs)
chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", plain_text))
print(f"Saved: {OUTPUT}")
print(f"Chinese characters in prose paragraphs: {chinese_chars}")
print(f"Tables: {len(doc.tables)}; paragraphs: {len(doc.paragraphs)}")
